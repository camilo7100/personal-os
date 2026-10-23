"""Extractos bancarios: lectura automática por banco, categorización y conciliación.

Lee los formatos reales de las cuentas del usuario (Bancolombia ahorros .xlsx, Fiducuenta .xlsx, RappiCard y RappiCuenta .pdf, tarjeta Nu .pdf),
abre los PDF protegidos con la clave guardada en el Administrador de credenciales (`statements-pdf-password`), importa los movimientos sin duplicar,
registra saldos (patrimonio), datos de las tarjetas (pago mínimo, fecha límite, cupo, cuotas) y qué meses están cubiertos por cuenta.
Regla de oro: lo que se mueve entre tus propias cuentas (traslados al fondo, recargas, pagos de tarjeta) NO es ingreso ni gasto.

Fuentes: archivos que el usuario deja en Documents\\PERSONAL_OS-Finanzas\\Extractos y adjuntos que llegan a su correo (buzón local de Thunderbird)."""
import email
import hashlib
import json
import os
import re
import shutil
from datetime import date, datetime, timedelta
from email import policy
from pathlib import Path

import typer

from . import db, money

app = typer.Typer(help="Extractos bancarios: importación automática por banco", no_args_is_help=True)
INBOX_DIR = Path(db.HOME_DIR) / "Documents" / "PERSONAL_OS-Finanzas" / "Extractos"
ARCHIVE = db.DATA_DIR / "finance" / "extractos"
HOLD_DIR = db.DATA_DIR / "finance" / "recibidos"  # lo que llega por Telegram: se guarda tal cual y NO se importa hasta que se revise
SENDERS = ("nu.com.co", "rappipay.co", "rappicard.co")
TRANSFER = ("Movimiento entre mis cuentas", "Pago de tarjeta de crédito")
MONTHS = {"ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "sep": 9, "sept": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12,
          "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}

# Categorías por defecto (en orden: la primera que coincide gana). config\money_rules.json puede agregar reglas propias, que van primero.
RULES: list[tuple[str, str]] = [
    (r"traslado (a|de) fondo de invers|retiro trasla|^adicion$|recarga \| pse|transferencia bancaria|recarga desde", "Movimiento entre mis cuentas"),
    (r"pago pse banco davivienda|pago pse nu compania|pago de rappicard|pagos por pse|pagos rappipay|gracias por tu pago|pago tarjeta nu|^pago$", "Pago de tarjeta de crédito"),
    (r"ajuste a tu favor|devoluci|reembolso|reverso", "Devoluciones y reembolsos"),
    (r"seguros del estado|soat|\bcda\b|tecnomec|revision tecnic|revisión técnic|taller|lavadero", "Vehículo (SOAT, revisión, taller)"),
    (r"veterinari|pet ?shop|mascota", "Mascotas"),
    (r"dhl|servientrega|interrapidisimo|coordinadora|envia\b", "Envíos"),
    (r"pago de nomi|nomina|nómina", "Salario"),
    (r"abono intereses|intereses ganados|rendimientos", "Intereses y rendimientos"),
    (r"cashback", "Cashback"),
    (r"transf de maria|envio de juan|envío de juan", "Familia (recibido)"),
    (r"transferencias a nequi", "Gastos vía Nequi"),
    (r"global colombia|global66|transferencia cta suc virtual|consig nacional efectivo", "Por aclarar"),
    (r"retefte|retencion|comision por cambio|comisión por cambio|cargos por conversion|cuota de manejo|gmf|4x1000|iva", "Impuestos y comisiones"),
    (r"anthropic|claude|openai|chatgpt|dlocal|google \*|apple\.com|netflix|spotify|disney|hbo|prime video|youtube", "Suscripciones y software"),
    (r"movistar|claro|tigo|wom|etb|epm|vanti|esp\b", "Servicios y telefonía"),
    (r"mercadoli|mercado pago|aliexpress|amazon|temu|shein|falabella\.com|mep\*", "Compras online"),
    (r"tienda d1|\bd1\b|ara\b|makro|exito|éxito|olimpica|olímpica|carulla|jumbo|market|supermerc|surtimax|avila market|oxxo", "Mercado"),
    (r"eds |gasolin|terpel|primax|texaco|biomax|peaje|uber|didi|cabify|indriver|parqueadero", "Transporte y gasolina"),
    (r"farmatodo|drogueria|droguería|farmacia|cruz verde|clinica|clínica|odontolog|laborator", "Salud y farmacia"),
    (r"wingo|avianca|latam|viva air|jetsmart|satena|copa air|booking|airbnb|hotel|despegar|decolar", "Viajes"),
    (r"cine colombia|cinemark|procinal|multiplex|pub\b|bar\b|beer|cervec|la jugada|gloriosa|discoteca|concierto|tuboleta|eticket|quick ticket", "Entretenimiento y salidas"),
    (r"arepas|parrill|ricco food|rest\b|restaurante|cocina|sancocho|gelat|helad|pizza|burger|hamburg|frisby|kfc|mcdonald|crepes|rappi\b|ifood|cafe|café|panader|donde jean|manhattan", "Comida fuera"),
    (r"barber|peluquer|spa\b|tatto|tattoo|uñas", "Cuidado personal"),
    (r"ropa|zenda|zara|koaj|arturo calle|adidas|nike", "Ropa"),
    (r"canchas|gym|gimnasio|smartfit|bodytech|deport", "Deporte"),
    (r"transf a |envio a otra llave|envío a otra llave|pago llave|transferencia cta suc virtual", "Transferencias a personas"),
    (r"deposito desde llave|depósito desde llave", "Otros ingresos"),
    (r"pago qr", "Compras varias (QR)"),
    (r"\bpse\b", "Pagos PSE (por aclarar)"),
]


def categorize(desc: str, amount: float) -> str:
    d = desc.lower()
    custom = money.load_rules()
    for cat, words in custom.items():
        if any(w.lower() in d for w in words):
            return cat
    for pat, cat in RULES:
        if re.search(pat, d):
            return cat
    return "Otros ingresos" if amount > 0 else "Sin categoría"


def is_transfer(category: str | None) -> bool:
    return (category or "") in TRANSFER


# ---------- utilidades de lectura ----------
def _us(s: str) -> float:
    """$1,234.56 / -1,234.56 (formato de RappiCuenta y Bancolombia en Excel)."""
    s = s.replace("$", "").replace(" ", "")
    neg = "-" in s
    v = float(re.sub(r"[^\d.]", "", s.replace(",", "")) or 0)
    return -v if neg else v


def _co(s: str) -> float:
    """$1.234,56 / -1.234,56 (formato de Nu, RappiCard y Fiducuenta)."""
    s = s.replace("$", "").replace(" ", "")
    neg = "-" in s
    v = float(re.sub(r"[^\d,]", "", s).replace(",", ".") or 0)
    return -v if neg else v


def pdf_passwords() -> list[str]:
    """Claves guardadas para abrir PDF (la del usuario y las extra, p. ej. la de la cuenta de su familiar): `statements-pdf-password`, `-2`, `-3`…"""
    import keyring
    out = []
    for key in ["statements-pdf-password"] + [f"statements-pdf-password-{i}" for i in range(2, 9)]:
        v = keyring.get_password("pos", key)
        if v:
            out.append(v)
    return out


def _pdf_text(path: Path) -> str:
    import fitz
    doc = fitz.open(path)
    if doc.needs_pass and not any(doc.authenticate(p) for p in pdf_passwords()):
        raise ValueError("PDF protegido: ninguna clave guardada lo abre (agrégala con `pos statements password --extra`)")
    return "\n".join(p.get_text() for p in doc)


def _ymd(day: int, mon: str, year: int) -> str:
    m = mon.lower().strip(".")
    return date(year, MONTHS.get(m) or MONTHS[m[:3]], day).isoformat()


# ---------- lectores por banco ----------
def parse_bancolombia_xlsx(path: Path) -> dict:
    import openpyxl
    rows = [[("" if c is None else str(c)).strip() for c in r] for ws in openpyxl.load_workbook(path, data_only=True).worksheets for r in ws.iter_rows(values_only=True)]
    info = next(rows[i + 1] for i, r in enumerate(rows) if r[:2] == ["DESDE", "HASTA"])
    start, end = (date(*map(int, info[0].split("/"))), date(*map(int, info[1].split("/"))))
    summ = next(rows[i + 1] for i, r in enumerate(rows) if r[:1] == ["SALDO ANTERIOR"])
    txs = []
    for r in rows:
        m = re.fullmatch(r"(\d{1,2})/(\d{1,2})", r[0] or "")
        if not m or not r[4]:
            continue
        d, mo = int(m.group(1)), int(m.group(2))
        y = end.year if mo <= end.month else end.year - 1 if start.year < end.year else end.year
        txs.append((date(y, mo, d).isoformat(), re.sub(r"\s+", " ", r[1].replace("\xa0", " ")).strip(), _us(r[4])))
    return {"bank": "bancolombia", "account": "Bancolombia ahorros", "kind": "savings", "number": info[3], "start": start.isoformat(), "end": end.isoformat(),
            "opening": _us(summ[0]), "closing": _us(summ[3]), "transactions": txs}


def parse_fiducuenta_xlsx(path: Path) -> dict:
    import openpyxl
    rows = [[("" if c is None else str(c)).strip() for c in r] for ws in openpyxl.load_workbook(path, data_only=True).worksheets for r in ws.iter_rows(values_only=True)]
    info = next(rows[i + 1] for i, r in enumerate(rows) if r[:2] == ["FONDO", "NIT"])
    start, end = (date(int(info[4][:4]), int(info[4][4:6]), int(info[4][6:])), date(int(info[5][:4]), int(info[5][4:6]), int(info[5][6:])))
    txs = []
    for r in rows:
        if re.fullmatch(r"\d{8}", r[0] or "") and r[2]:
            kind = r[1].upper()
            amt = _co(r[2])
            amt = amt if kind.startswith("ADICION") else -amt
            txs.append((f"{r[0][:4]}-{r[0][4:6]}-{r[0][6:]}", kind.title() if not kind.startswith("RETEFTE") else "Retefte Fiducuenta", amt))
    s = next(rows[i + 1] for i, r in enumerate(rows) if r[:2] == ["VALOR EN PESOS", "VALOR EN UNIDADES"])
    yields = _co(s[4])
    if yields:
        txs.append((end.isoformat(), "Rendimientos netos Fiducuenta", yields))
    return {"bank": "fiducuenta", "account": "Fiducuenta (inversión)", "kind": "invest", "number": info[3], "start": start.isoformat(), "end": end.isoformat(),
            "opening": _co(s[0]), "closing": _co(s[6]), "transactions": txs, "rate": info[7]}


def parse_fiducuenta_pdf(text: str) -> dict:
    """El mismo extracto de la Fiducuenta pero en PDF (cada valor en su propia línea). Produce los mismos movimientos que el .xlsx, así que reimportar no duplica."""
    d8 = lambda s: f"{s[:4]}-{s[4:6]}-{s[6:]}"
    start = d8(re.search(r"Desde:\s*(\d{8})", text).group(1))
    end = d8(re.search(r"Hasta:\s*(\d{8})", text).group(1))
    summary = text[text.index("SALDO ANTERIOR"):text.index("FECHA.")]
    n = [_co(x) for x in re.findall(r"\d{1,3}(?:\.\d{3})*,\d+", summary)]  # saldo anterior, unidades, adiciones, retiros, rendimientos, retención, nuevo saldo, unidades
    lines = [l.strip() for l in text[text.index("FECHA."):].splitlines()]  # solo la tabla de movimientos (el encabezado también trae fechas de 8 dígitos)
    txs = []
    for i, l in enumerate(lines):
        if re.fullmatch(r"\d{8}", l) and i + 2 < len(lines) and re.fullmatch(r"\d{1,3}(?:\.\d{3})*,\d+", lines[i + 2]):
            kind, amt = lines[i + 1].upper(), _co(lines[i + 2])
            txs.append((d8(l), kind.title() if not kind.startswith("RETEFTE") else "Retefte Fiducuenta", amt if kind.startswith("ADICION") else -amt))
    if n[4]:
        txs.append((end, "Rendimientos netos Fiducuenta", n[4]))
    num = re.search(r"Cuenta de Inversi\w+:\s*(\d+)", text)
    return {"bank": "fiducuenta", "account": "Fiducuenta (inversión)", "kind": "invest", "number": num.group(1) if num else "", "start": start, "end": end,
            "opening": n[0], "closing": n[6], "transactions": txs, "rate": (re.search(r"Rentabilidad Periodo:\s*([\d,]+ ?% ?\w*)", text) or [None, ""])[1]}


POCKET = "Bolsillo propio (cuenta de un familiar)"


def parse_family_rappicuenta(text: str, st: dict) -> dict:
    """Extracto de la RappiCuenta de un familiar donde tienes un bolsillo. Sus movimientos NO son tuyos: no entran a tus ingresos ni gastos.
    Solo se guardan los totales (para estimar el rendimiento de tu bolsillo) y el periodo cubierto."""
    g = lambda pat: _us(re.search(pat, text).group(1)) if re.search(pat, text) else 0.0
    interest, retention = g(r"Intereses ganados\s*\n?\s*([+-]?\$[\d,.]+)"), abs(g(r"Retención en la fuente\s*\n?\s*([+-]?\$[\d,.]+)"))
    return {"bank": "rappicuenta-papa", "account": POCKET, "kind": "savings", "third_party": True, "holder": st["holder"], "number": st["number"],
            "start": st["start"], "end": st["end"], "opening": st["opening"], "closing": st["closing"], "transactions": [], "interest_net": round(interest - retention, 2)}


def parse_rappicard_pdf(text: str) -> dict:
    g = lambda pat: re.search(pat, text, re.S)
    p = g(r"Periodo facturado\s*Desde\s*(\d{1,2}) (\w+) (\d{4})\s*Hasta\s*(\d{1,2}) (\w+) (\d{4})\s*Fecha de pago\s*(\d{1,2}) (\w+) (\d{4})")
    start, end, due = _ymd(int(p.group(1)), p.group(2), int(p.group(3))), _ymd(int(p.group(4)), p.group(5), int(p.group(6))), _ymd(int(p.group(7)), p.group(8), int(p.group(9)))
    total = _co(g(r"Pago total\s*\$\s*([\d.,]+)").group(1))
    minimum = _co(g(r"Pago mínimo\s*\$\s*([\d.,]+)").group(1))
    limit_ = _co(g(r"Cupo total:\s*\$([\d.,]+)").group(1))
    lines = [l.strip() for l in text.splitlines()]
    txs, pending = [], 0.0
    for i, l in enumerate(lines):
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", l):
            continue
        j, desc = i + 1, []
        while j < len(lines) and not lines[j].startswith("$"):
            desc.append(lines[j])
            j += 1
        if j >= len(lines):
            continue
        val = _co(lines[j])
        d = " ".join(desc)
        txs.append((l, d, -val))  # compra = sale plata (negativo); pago a la tarjeta (valor negativo en el extracto) = entra (positivo)
        if val > 0 and j + 3 < len(lines) and lines[j + 3].startswith("$"):
            pending += _co(lines[j + 3])
    return {"bank": "rappicard", "account": "RappiCard (crédito)", "kind": "credit", "start": start, "end": end, "closing": -total, "transactions": txs,
            "card": {"due": due, "minimum": minimum, "total": total, "limit": limit_, "used": total, "pending_installments": round(pending, 2)}}


def parse_rappicuenta_pdf(text: str) -> dict:
    gen = re.search(r"generado el (\d{1,2}) de (\w+) de (\d{4})", text)
    gen_d = date(int(gen.group(3)), MONTHS[gen.group(2).lower()], int(gen.group(1)))
    end = gen_d - timedelta(days=1)
    start = end.replace(day=1)
    opening = _us(re.search(r"Saldo anterior\s*(\$[\d,.]+)", text).group(1))
    closing = _us(re.search(r"Saldo final\s*(\$[\d,.]+)", text).group(1))
    lines = [l.strip() for l in text.splitlines()]
    txs = []
    for i, l in enumerate(lines):
        m = re.fullmatch(r"(\d{1,2}) (\w{3}) (\d{4})", l)
        if m and i + 2 < len(lines) and re.match(r"[+-]\$", lines[i + 2]):
            txs.append((_ymd(int(m.group(1)), m.group(2), int(m.group(3))), lines[i + 1], _us(lines[i + 2])))
    holder = re.search(r"Extracto de tu RappiCuenta\s*\n\s*(.+)", text)
    number = re.search(r"Número de cuenta:\s*\n?\s*(\d+)", text)
    return {"holder": holder.group(1).strip() if holder else "", "number": number.group(1) if number else "",
            "bank": "rappicuenta", "account": "RappiCuenta (ahorros)", "kind": "savings", "start": start.isoformat(), "end": end.isoformat(),
            "opening": opening, "closing": closing, "transactions": txs}


def parse_nu_card_pdf(text: str) -> dict:
    corte = re.search(r"Fecha de corte\s*(\d{1,2}) (\w{3}) (\d{4})", text)
    end = date(int(corte.group(3)), MONTHS[corte.group(2).lower()], int(corte.group(1))) - timedelta(days=1)
    lim = re.search(r"Fecha límite de pago\s*(\d{1,2}) (\w{3}) (\d{4})", text)
    due = _ymd(int(lim.group(1)), lim.group(2), int(lim.group(3)))
    # deuda total del corte (en los extractos reenviados, «Usado» muestra el saldo del día del reenvío, no el del corte)
    tot = re.search(r"\$([\d.,]+)\s*(?:PAGO|DEUDA TOTAL) HASTA", text)
    used = _co(tot.group(1)) if tot else _co(re.search(r"Usado\s*\$([\d.,]+)", text).group(1))
    limit_ = _co(re.search(r"Tu cupo definido\s*\$([\d.,]+)", text).group(1))
    minimum = _co(re.search(r"PAGO MÍNIMO\s*\$([\d.,]+)", text).group(1))
    lines = [l.strip() for l in text.splitlines()]
    txs = []
    i = lines.index("Restante") if "Restante" in lines else 0  # la tabla de movimientos empieza tras su encabezado
    while i < len(lines):
        m = re.fullmatch(r"(\d{1,2}) (ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SEP|OCT|NOV|DIC)(?: (\d{4}))?", lines[i])
        if not m:
            i += 1
            continue
        j = i + 1
        year = m.group(3)
        if not year and j < len(lines) and re.fullmatch(r"\d{4}", lines[j]):
            year, j = lines[j], j + 1
        if not year:
            i += 1
            continue
        desc = []
        while j < len(lines) and not lines[j].startswith("$") and len(desc) < 4:
            desc.append(lines[j])
            j += 1
        if j >= len(lines) or not lines[j].startswith("$"):
            i += 1
            continue
        d = " ".join(desc).strip()
        day = _ymd(int(m.group(1)), m.group(2), int(year))
        val = _co(lines[j])
        if re.search(r"pago", d, re.I):
            txs.append((day, "Pago tarjeta Nu", val))
        elif re.search(r"devoluci|ajuste a tu favor|reembolso", d, re.I):
            txs.append((day, d, abs(val)))  # te devolvieron plata: entra
        else:
            txs.append((day, d, -val))
        k = j + 1
        while k < len(lines) and not re.fullmatch(r"(\d{1,2}) (ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SEP|OCT|NOV|DIC)(?: \d{4})?", lines[k]) and not lines[k].startswith("Pago mínimo"):
            if lines[k].startswith("↪ Comisión por cambio") and k + 1 < len(lines) and lines[k + 1].startswith("$"):
                txs.append((day, f"Comisión por cambio de moneda ({d[:30]})", -_co(lines[k + 1])))
            k += 1
        i = k
    return {"bank": "nu", "account": "Nu (crédito)", "kind": "credit", "start": (end - timedelta(days=30)).isoformat(), "end": end.isoformat(), "closing": -used,
            "transactions": txs, "card": {"due": due, "minimum": minimum, "total": used, "limit": limit_, "used": used, "pending_installments": 0.0}}


def read(path: Path) -> dict | None:
    """Detecta de qué banco es el archivo y lo lee. None si no es un extracto conocido (o es una cuenta inactiva)."""
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        head = " ".join(str(c) for ws in openpyxl.load_workbook(path, data_only=True).worksheets for r in ws.iter_rows(max_row=8, values_only=True) for c in r if c)
        if "FIDUCUENTA" in head.upper():
            return parse_fiducuenta_xlsx(path)
        if "CUENTA DE AHORROS" in head.upper() or "SALDO ANTERIOR" in head.upper() or "DESDE" in head.upper():
            return parse_bancolombia_xlsx(path)
        return None
    if path.suffix.lower() == ".pdf":
        text = _pdf_text(path)
        if "FIDUCUENTA FIDUCIARIA BANCOLOMBIA" in text.upper() and "EXTRACTO MENSUAL" in text.upper():
            return parse_fiducuenta_pdf(text)
        if "Extracto de tarjeta de crédito" in text and ("Rappi" in text or "Davivienda" in text):
            return parse_rappicard_pdf(text)
        if "Extracto de tu RappiCuenta" in text:
            st = parse_rappicuenta_pdf(text)
            if st["holder"] and st["holder"].lower() != money.assume_config()["owner_name"]:
                return parse_family_rappicuenta(text, st)  # cuenta de otra persona (p. ej. tu familiar): solo cuenta tu bolsillo
            return st
        if "Nu Financiera" in text and "Tus tarjetas" in text:
            return parse_nu_card_pdf(text)
    return None


# ---------- importación ----------
def _account(con, name: str, kind: str) -> int:
    con.execute("INSERT OR IGNORE INTO accounts(name, kind) VALUES (?,?)", (name, kind))
    return con.execute("SELECT id FROM accounts WHERE name=?", (name,)).fetchone()["id"]


def import_parsed(con, st: dict, source: str = "") -> dict:
    acc = _account(con, st["account"], st["kind"])
    if st.get("third_party"):  # cuenta de un familiar: solo periodo cubierto + totales para estimar tu bolsillo (sin movimientos ni saldo propio)
        con.execute("INSERT OR REPLACE INTO statement_periods(account_id, start, end, source) VALUES (?,?,?,?)", (acc, st["start"], st["end"], source))
        con.execute("INSERT OR REPLACE INTO meta(key, value) VALUES (?, ?)", (f"family:rappicuenta:{st['end'][:7]}", json.dumps(
            {"opening": st["opening"], "closing": st["closing"], "interest_net": st["interest_net"], "number": st["number"]})))
        db.audit(con, "statements.import", "WRITE", st["account"], detail=f"{st['start']}..{st['end']} · cuenta de {st['holder']} (solo totales)", actor="job")
        from . import finance
        finance.refresh_estimates(con)
        return {"account": st["account"], "period": f"{st['start']} → {st['end']}", "added": 0, "total": 0, "closing": st["closing"], "family": True}
    added = 0
    seen: dict[str, int] = {}
    for d, desc, amt in st["transactions"]:
        key = f"{acc}|{d}|{desc}|{round(amt, 2)}"
        seen[key] = seen.get(key, 0) + 1
        h = hashlib.sha1(f"{key}|{seen[key]}".encode()).hexdigest()
        cur = con.execute("INSERT OR IGNORE INTO transactions(account_id, date, description, amount, category, hash) VALUES (?,?,?,?,?,?)",
                          (acc, d, desc, round(amt, 2), categorize(desc, amt), h))
        added += cur.rowcount
    con.execute("INSERT INTO balances(account_id, date, amount) VALUES (?,?,?) ON CONFLICT(account_id, date) DO UPDATE SET amount=excluded.amount, estimated=0", (acc, st["end"], round(st["closing"], 2)))  # el dato real reemplaza al supuesto
    if st.get("opening") is not None:
        prev = (date.fromisoformat(st["start"]) - timedelta(days=1)).isoformat()
        con.execute("INSERT INTO balances(account_id, date, amount) VALUES (?,?,?) ON CONFLICT(account_id, date) DO UPDATE SET amount=excluded.amount, estimated=0 WHERE estimated=1", (acc, prev, round(st["opening"], 2)))
    con.execute("INSERT OR REPLACE INTO statement_periods(account_id, start, end, source) VALUES (?,?,?,?)", (acc, st["start"], st["end"], source))
    if st.get("card"):
        c = st["card"]
        con.execute("INSERT OR REPLACE INTO card_statements(account_id, period_end, due_date, min_payment, total_due, credit_limit, used, pending_installments) VALUES (?,?,?,?,?,?,?,?)",
                    (acc, st["end"], c["due"], c["minimum"], c["total"], c["limit"], c["used"], c["pending_installments"]))
        _update_payment_task(con, st["account"], c)
    db.audit(con, "statements.import", "WRITE", st["account"], detail=f"{st['start']}..{st['end']} · {added} movimientos nuevos", actor="job")
    return {"account": st["account"], "period": f"{st['start']} → {st['end']}", "added": added, "total": len(st["transactions"]), "closing": st["closing"]}


def _update_payment_task(con, account: str, c: dict) -> None:
    """La tarea mensual de pago de la tarjeta queda con el monto y la fecha reales del último extracto."""
    pat = "%Nu%" if account.startswith("Nu") else "%Rappi%"
    due = date.fromisoformat(c["due"])
    if due < date.today():
        return  # extracto viejo: no toca la tarea del mes en curso
    for t in con.execute("SELECT id, due_date FROM tasks WHERE status IN ('todo','doing','blocked') AND tags LIKE '%pago%' AND title LIKE ?", (pat,)).fetchall():
        note = (f"Último extracto: pago total ${c['total']:,.0f} · pago mínimo ${c['minimum']:,.0f} · fecha límite {c['due']}. "
                f"Paga el TOTAL para no pagar intereses.").replace(",", ".")
        new_due = (due - timedelta(days=1)).isoformat() if t["due_date"] and abs((date.fromisoformat(t["due_date"]) - due).days) <= 20 else t["due_date"]
        con.execute("UPDATE tasks SET notes=?, due_date=? WHERE id=?", (note, new_due, t["id"]))


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def import_file(con, path: Path, source: str = "manual") -> dict | None:
    h = file_hash(path)
    if con.execute("SELECT 1 FROM statement_files WHERE hash=? AND account IS NOT NULL", (h,)).fetchone():
        return {"skipped": path.name}
    st = read(path)
    if not st:
        con.execute("INSERT OR REPLACE INTO statement_files(hash, name, account, period, source) VALUES (?,?,?,?,?)", (h, path.name, None, None, source))
        return None
    r = import_parsed(con, st, source)
    dst = ARCHIVE / st["bank"] / f"{st['end'][:7]}{path.suffix.lower()}"
    dst.parent.mkdir(parents=True, exist_ok=True)
    if not dst.exists():
        shutil.copy2(path, dst)
    con.execute("INSERT OR REPLACE INTO statement_files(hash, name, account, period, source) VALUES (?,?,?,?,?)", (h, path.name, st["account"], st["end"][:7], source))
    return r


def scan_mailbox(con) -> list[dict]:
    """Adjuntos de extractos que llegaron al correo (buzón local de Thunderbird, solo lectura)."""
    from . import mail
    prof = mail._find_profile()
    out = []
    if not prof:
        return out
    tmp = db.DATA_DIR / "finance" / "_entrada"
    tmp.mkdir(parents=True, exist_ok=True)
    for dirpath, _, files in os.walk(Path(prof) / "ImapMail"):
        if not (dirpath.endswith("cur") or dirpath.endswith("new")):
            continue
        for f in files:
            p = Path(dirpath) / f
            try:
                raw = p.read_bytes()
            except OSError:
                continue
            if not any(s.encode() in raw[:8000] for s in SENDERS):
                continue
            msg = email.message_from_bytes(raw, policy=policy.default)
            if not any(s in str(msg.get("From", "")) for s in SENDERS):
                continue
            for part in msg.walk():
                fn = part.get_filename()
                if fn and fn.lower().endswith(".pdf"):
                    data = part.get_payload(decode=True)
                    if con.execute("SELECT 1 FROM statement_files WHERE hash=?", (hashlib.sha256(data).hexdigest(),)).fetchone():
                        continue
                    q = tmp / re.sub(r"[^\w.\-]", "_", fn)
                    q.write_bytes(data)
                    try:
                        r = import_file(con, q, "correo")
                        if r and "account" in r:
                            out.append(r)
                    except Exception as e:  # un extracto raro no debe tumbar el resto
                        out.append({"error": f"{fn}: {e}"})
                    finally:
                        q.unlink(missing_ok=True)
    return out


def scan_folder(con, folder: Path | None = None) -> list[dict]:
    folder = folder or INBOX_DIR
    folder.mkdir(parents=True, exist_ok=True)
    out = []
    for p in sorted(folder.iterdir()):
        if p.suffix.lower() in (".pdf", ".xlsx") and p.is_file():
            try:
                r = import_file(con, p, "carpeta")
                if r and "account" in r:
                    out.append(r)
                if r:  # importado (o ya estaba): queda archivado en data\finance\extractos, la carpeta de entrada vuelve a quedar vacía
                    p.unlink(missing_ok=True)
            except Exception as e:
                out.append({"error": f"{p.name}: {e}"})
    return out


def coverage(con, months: int = 6, today: date | None = None) -> dict:
    """Qué meses tiene cada cuenta cubiertos por extractos (para saber si un mes está completo)."""
    today = today or date.today()
    ms = []
    y, m = today.year, today.month
    for _ in range(months):
        ms.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    ms.reverse()
    rows = []
    for a in con.execute("SELECT id, name, kind FROM accounts WHERE id IN (SELECT account_id FROM statement_periods) ORDER BY name"):
        per = con.execute("SELECT start, end FROM statement_periods WHERE account_id=?", (a["id"],)).fetchall()
        cells = []
        for mo in ms:
            y_, m_ = map(int, mo.split("-"))
            first = date(y_, m_, 1)
            last = (first + timedelta(days=32)).replace(day=1) - timedelta(days=1)
            # un mes cuenta como cubierto si algún extracto abarca al menos 20 de sus días
            covered = any((min(date.fromisoformat(p["end"]), last) - max(date.fromisoformat(p["start"]), first)).days + 1 >= 20 for p in per)
            cells.append(covered)
        rows.append({"account": a["name"], "kind": a["kind"], "cells": cells, "assumed": money.is_assumed_account(a["name"])})
    # una cuenta «supuesta» (Bancolombia: ingreso supuesto hasta que llegue su extracto) no bloquea que el mes se dé por completo
    for r in rows:
        r["lag"] = money.is_lag_account(r["account"])
    firm = [r for r in rows if not r["lag"]] or rows  # si solo existe esa cuenta, ella decide
    complete = [all(r["cells"][i] for r in firm) if firm else False for i in range(len(ms))]
    return {"months": ms, "rows": rows, "complete": complete}


# ---------- archivos que llegan por Telegram: se guardan y esperan revisión ----------
def hold_file(src: Path, original_name: str) -> Path:
    """Guarda un archivo recibido por Telegram sin tocarlo ni importarlo (los formatos de los bancos son feos: se revisan a mano y se suben como convenga)."""
    HOLD_DIR.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w.\- ]", "_", original_name)[:120] or "archivo"
    dest = HOLD_DIR / f"{datetime.now():%Y%m%d-%H%M%S}_{safe}"
    shutil.move(str(src), dest)
    return dest


def held() -> list[Path]:
    return sorted(HOLD_DIR.glob("*")) if HOLD_DIR.exists() else []


# ---------- el día 15: tarea de subir los extractos y conteo uno por uno ----------
UPLOAD_DAY = 15
MES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def prev_month(today: date) -> str:
    return (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")


def _month_label(ym: str) -> str:
    return f"{MES[int(ym[5:])]} {ym[:4]}"


def covered(con, account_id: int, month: str) -> bool:
    """¿Algún extracto de esta cuenta abarca al menos 20 días de ese mes?"""
    y, m = map(int, month.split("-"))
    first = date(y, m, 1)
    last = (first + timedelta(days=32)).replace(day=1) - timedelta(days=1)
    return any((min(date.fromisoformat(p["end"]), last) - max(date.fromisoformat(p["start"]), first)).days + 1 >= 20
               for p in con.execute("SELECT start, end FROM statement_periods WHERE account_id=?", (account_id,)))


def expected_accounts(con) -> list[dict]:
    """Cuentas de las que se esperan extractos (Nequi no tiene: es una deuda estimada)."""
    return [dict(r) for r in con.execute("SELECT id, name FROM accounts WHERE name NOT LIKE 'Nequi%' ORDER BY name")]


def upload_task(con, month: str) -> int | None:
    r = con.execute("SELECT task_id FROM external_items WHERE uid=?", (f"statements:{month}",)).fetchone()
    return r["task_id"] if r else None


def checklist(con, month: str) -> list[dict]:
    """Un renglón por cuenta: ¿ya llegó el extracto de ese mes?"""
    return [{"account": a["name"], "id": a["id"], "done": covered(con, a["id"], month)} for a in expected_accounts(con)]


def ensure_upload_task(con, today: date | None = None) -> int | None:
    """El día 15 (o después, si el PC estuvo apagado) crea la tarea «Subir los extractos de <mes anterior>» con un subpaso por cuenta. Una vez por mes."""
    from . import services
    today = today or date.today()
    month = prev_month(today)
    if today.day < UPLOAD_DAY or upload_task(con, month) or not expected_accounts(con):
        return None
    due = today.replace(day=UPLOAD_DAY)
    cur = con.execute("INSERT INTO tasks(title, notes, priority, due_date, source, tags, est_minutes) VALUES (?,?,?,?,?,?,?)",
                      (f"📄 Subir los extractos de {_month_label(month)}",
                       "Baja el extracto de cada cuenta y suéltalo en Documents\\PERSONAL_OS-Finanzas\\Extractos (o mándalo a tu correo). Yo los cuento uno por uno, los reviso y los cuadro.",
                       2, due.isoformat(), "auto", "finanzas,extractos", 30))
    tid = cur.lastrowid
    con.execute("INSERT INTO external_items(uid, source, task_id, due) VALUES (?,?,?,?)", (f"statements:{month}", "auto", tid, due.isoformat()))
    for c in checklist(con, month):
        services.add_subtask(con, tid, f"Extracto de {c['account']}", actor="job")
    db.audit(con, "autotask", "WRITE", f"task:{tid}", detail=f"extractos {month}", actor="job")
    return tid


def tick_checklist(con, today: date | None = None) -> dict | None:
    """Marca como hecho el subpaso de cada cuenta cuyo extracto ya llegó y cierra la tarea cuando están todos. Devuelve el avance."""
    from . import services
    today = today or date.today()
    month = prev_month(today)
    tid = upload_task(con, month)
    if not tid:
        return None
    items = checklist(con, month)
    subs = {s["title"]: s for s in services.subtasks(con, tid)}
    for c in items:
        s = subs.get(f"Extracto de {c['account']}")
        if c["done"] and s and s["status"] != "done":
            services.done_task(con, s["id"], actor="job")
    got = [c["account"] for c in items if c["done"]]
    missing = [c["account"] for c in items if not c["done"]]
    if not missing:
        services.done_task(con, tid, actor="job")
    return {"month": month, "got": got, "missing": missing, "total": len(items)}


def progress_message(p: dict) -> str:
    head = f"📥 Extractos de {_month_label(p['month'])}: {len(p['got'])}/{p['total']}"
    if not p["missing"]:
        return head + " ✅ ¡Todos llegaron! Los reviso y los cuadro en el panel de Finanzas."
    wait = f"\n📎 {len(held())} archivo(s) tuyos esperan mi revisión (no se importan solos)." if held() else ""
    return head + wait + f"\nFaltan: " + ", ".join(p["missing"]) + f"\nSuéltalos en Documents\\PERSONAL_OS-Finanzas\\Extractos o mándalos a tu correo."


def run_job() -> str:
    """Job diario: importa extractos nuevos del correo y de la carpeta, y avisa por Telegram con lo importante."""
    from . import notify
    with db.session() as con:
        found = scan_mailbox(con) + scan_folder(con)
    ok = [r for r in found if "account" in r]
    if ok:
        lines = [f"📄 {r['account']}: {r['period']} · {r['added']} movimientos" for r in ok]
        with db.session() as con:
            for c in con.execute("SELECT a.name, c.due_date, c.total_due, c.min_payment FROM card_statements c JOIN accounts a ON a.id=c.account_id "
                                 "WHERE c.period_end=(SELECT MAX(period_end) FROM card_statements WHERE account_id=c.account_id) AND c.due_date >= date('now')"):
                lines.append(f"💳 {c['name']}: paga ${c['total_due']:,.0f} antes del {c['due_date']}".replace(",", "."))
        notify.push_text("Nuevos extractos importados\n" + "\n".join(lines) + "\nRevisa el panel: Finanzas.")
    from . import finance as _fin
    with db.session() as con:
        _fin.refresh_estimates(con)  # saldos supuestos mes a mes de las cuentas que se atrasan (Fiducuenta y tu bolsillo)
    # el día 15: tarea de subir extractos y conteo uno por uno (aviso solo cuando cambia el avance)
    today = date.today()
    prev = prev_month(today)
    with db.session() as con:
        created = ensure_upload_task(con, today)
        prog = tick_checklist(con, today)
        if prog:
            key, state = f"statements:progress:{prev}", f"{len(prog['got'])}/{prog['total']}"
            last = con.execute("SELECT value FROM meta WHERE key=?", (key,)).fetchone()
            if created or not last or last["value"] != state:
                notify.push_text(progress_message(prog))
                con.execute("INSERT OR REPLACE INTO meta(key, value) VALUES (?, ?)", (key, state))
            if not prog["missing"]:  # todos llegaron: queda pendiente la revisión que hace Claude (skill finance-review)
                con.execute("INSERT OR IGNORE INTO meta(key, value) VALUES (?, 'pending')", (f"statements:review:{prev}",))
    # resumen del mes al celular (una sola vez): cuando llegaron todos los extractos, o el día 20 con lo que haya
    from . import finance
    with db.session() as con:
        key = f"finance:report:{prev}"
        all_in = con.execute("SELECT 1 FROM meta WHERE key=?", (f"statements:review:{prev}",)).fetchone()
        if (all_in or today.day >= 20) and not con.execute("SELECT 1 FROM meta WHERE key=?", (key,)).fetchone() and con.execute("SELECT 1 FROM transactions LIMIT 1").fetchone():
            notify.push_text(finance.monthly_report(con, prev))
            con.execute("INSERT OR REPLACE INTO meta(key, value) VALUES (?, '1')", (key,))
    errs = [r["error"] for r in found if "error" in r]
    return f"{len(ok)} extractos nuevos" + (f" · {len(errs)} con error: {'; '.join(errs)[:200]}" if errs else "")


# ---------- CLI ----------
@app.command("import")
def import_cmd(paths: list[Path] = typer.Argument(..., help="Archivos .pdf o .xlsx de extractos")) -> None:
    """Importa extractos (detecta el banco solo; reimportar no duplica)."""
    with db.session() as con:
        for p in paths:
            r = import_file(con, p)
            typer.echo(f"{p.name}: " + (json.dumps(r, ensure_ascii=False) if r else "no es un extracto conocido"))


@app.command("held")
def held_cmd() -> None:
    """Lista los archivos recibidos por Telegram que esperan revisión."""
    for p in held():
        typer.echo(f"{p.name}  ({p.stat().st_size // 1024} KB)")
    if not held():
        typer.echo("Nada en espera.")


@app.command("import-held")
def import_held_cmd() -> None:
    """Importa los archivos en espera que se reconozcan (los quita de la espera); los demás se quedan para revisarlos a mano."""
    with db.session() as con:
        for p in held():
            try:
                r = import_file(con, p)
            except Exception as e:  # formato raro: queda en espera
                typer.echo(f"{p.name}: no se pudo ({e})")
                continue
            if r:
                typer.echo(f"{p.name}: " + json.dumps(r, ensure_ascii=False))
                p.unlink(missing_ok=True)
            else:
                typer.echo(f"{p.name}: formato no reconocido, sigue en espera")


@app.command("scan")
def scan_cmd() -> None:
    """Busca extractos nuevos en el correo y en Documents\\PERSONAL_OS-Finanzas\\Extractos."""
    typer.echo(run_job())


@app.command("password")
def password_cmd(extra: bool = typer.Option(False, "--extra", help="Añade otra clave (se prueban todas) en vez de reemplazar la principal")) -> None:
    """Guarda la clave de los PDF de extractos en el Administrador de credenciales de Windows."""
    import keyring
    key = "statements-pdf-password"
    if extra:
        key = next(k for k in [f"statements-pdf-password-{i}" for i in range(2, 9)] if not keyring.get_password("pos", k))
    keyring.set_password("pos", key, typer.prompt("Clave de los PDF", hide_input=True))
    typer.echo("ok")


@app.command("recategorize")
def recategorize_cmd() -> None:
    """Vuelve a categorizar todos los movimientos con las reglas actuales (respeta lo que categorizaste a mano)."""
    with db.session() as con:
        n = 0
        for t in con.execute("SELECT id, description, amount, category FROM transactions WHERE manual_category=0").fetchall():
            c = categorize(t["description"], t["amount"])
            if c != t["category"]:
                con.execute("UPDATE transactions SET category=? WHERE id=?", (c, t["id"]))
                n += 1
    typer.echo(f"{n} movimientos recategorizados")
